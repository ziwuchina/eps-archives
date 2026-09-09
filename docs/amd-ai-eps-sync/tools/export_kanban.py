# -*- coding: utf-8 -*-
"""将裴谦-任务看板.xlsx 导出为 JSON 存档，放入 eps-archives 仓库"""
import openpyxl
import json
import re

SRC = r"C:\Users\Administrator\Desktop\AI相关\EPS项目汇报\裴谦-任务看板.xlsx"
OUT = r"C:\Users\Administrator\Desktop\AI相关\EPS项目汇报\eps-archives\reports\kanban-2026-09-04.json"

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["数据表"]
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1, max_col=11))]
records = []
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=11):
    cells = [c.value for c in row]
    if cells[0] is None or str(cells[0]).strip() == "":
        continue
    rec = {}
    for i, h in enumerate(headers):
        v = cells[i]
        if v is not None:
            s = str(v).strip()
            # 脱敏：替换密码与账号/密码组合
            s = s.replace("281436G52v364i.", "******")
            s = s.replace("曾玮/281436G52v364i.", "曾玮/******")
            rec[str(h)] = s
    records.append(rec)

with open(OUT, "w", encoding="utf-8") as f:
    json.dump({"source": "裴谦-任务看板.xlsx", "archived_at": "2026-09-04", "count": len(records), "records": records}, f, ensure_ascii=False, indent=2)

print("saved:", len(records), "records")
